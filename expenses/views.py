import csv
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.mail import send_mail
from django.core.paginator import Paginator
from django.db.models import Sum, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

# ⭐ Forms
from .forms import (
    ExpenseForm, SignUpForm, BudgetForm, SavingsGoalForm,
    RecurringExpenseForm, SplitExpenseForm, SplitParticipantFormSet,
    UserProfileForm, GroupForm, GroupExpenseForm, GroupInviteForm,
)

# ⭐ Models
from .models import (
    Expense, Budget, SavingsGoal, CATEGORY_CHOICES, CATEGORY_META,
    RecurringExpense, SplitExpense, SplitParticipant, UserProfile,
    CURRENCY_SYMBOLS, CURRENCY_CHOICES, EXCHANGE_RATES_TO_INR,
    Group, GroupMember, GroupExpense,
)

# ⭐ Utils & ML & SMS  — ITHU THAAN MISSING-A IRUNDHADHU!
from .utils import (
    predict_category,
    generate_insights,
    build_heatmap,
    analyze_personality,
)
from .ml_predictor import smart_predict, train_model, model_info
from .sms_utils import send_sms, check_budget_alerts


# ============================================================
# HELPERS
# ============================================================
def _total(qs):
    return qs.aggregate(t=Sum('amount_in_inr'))['t'] or Decimal('0')


def _apply_filters(request, qs):
    today = timezone.localdate()
    period = request.GET.get('period', 'all')

    if period == 'today':
        qs = qs.filter(date=today)
    elif period == 'week':
        qs = qs.filter(date__gte=today - timedelta(days=today.weekday()), date__lte=today)
    elif period == 'month':
        qs = qs.filter(date__gte=today.replace(day=1), date__lte=today)
    elif period == 'year':
        qs = qs.filter(date__gte=today.replace(month=1, day=1), date__lte=today)
    elif period == 'custom':
        start = request.GET.get('start')
        end = request.GET.get('end')
        if start:
            qs = qs.filter(date__gte=start)
        if end:
            qs = qs.filter(date__lte=end)

    category = request.GET.get('category', '').strip()
    if category:
        qs = qs.filter(category=category)

    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(title__icontains=q) | Q(notes__icontains=q))

    sort = request.GET.get('sort', '-date')
    allowed = {'-date', 'date', '-amount', 'amount', 'category', '-created_at'}
    if sort not in allowed:
        sort = '-date'

    return qs.order_by(sort, '-id'), period


# ============================================================
# AUTH
# ============================================================
def signup(request):
    if request.user.is_authenticated:
        return redirect('dashboard')

    if request.method == 'POST':
        form = SignUpForm(request.POST)
        if form.is_valid():
            user = form.save()
            UserProfile.objects.get_or_create(user=user)
            messages.success(
                request,
                f'Account created successfully, {user.username}! 🎉 Please sign in to continue.'
            )
            return redirect('login')
        messages.error(request, 'Please fix the errors below.')
    else:
        form = SignUpForm()

    return render(request, 'registration/signup.html', {'form': form})


# ============================================================
# DASHBOARD
# ============================================================
@login_required
def dashboard(request):
    user = request.user
    today = timezone.localdate()
    qs = Expense.objects.filter(user=user)

    week_start = today - timedelta(days=today.weekday())
    month_start = today.replace(day=1)
    year_start = today.replace(month=1, day=1)

    totals = {
        'today': _total(qs.filter(date=today)),
        'week': _total(qs.filter(date__gte=week_start, date__lte=today)),
        'month': _total(qs.filter(date__gte=month_start, date__lte=today)),
        'year': _total(qs.filter(date__gte=year_start, date__lte=today)),
        'all': _total(qs),
    }

    cat_qs = (qs.filter(date__gte=month_start)
                .values('category')
                .annotate(total=Sum('amount_in_inr'))
                .order_by('-total'))

    category_labels, category_values, category_colors = [], [], []
    for row in cat_qs:
        cat = row['category'] or 'Other'
        category_labels.append(cat)
        category_values.append(float(row['total']))
        category_colors.append(CATEGORY_META.get(cat, CATEGORY_META['Other'])['color'])

    top_category = category_labels[0] if category_labels else '—'

    trend_labels, trend_values = [], []
    for i in range(6, -1, -1):
        d = today - timedelta(days=i)
        trend_labels.append(d.strftime('%a %d'))
        trend_values.append(float(_total(qs.filter(date=d))))

    budget = Budget.objects.filter(user=user, month=month_start).first()
    budget_amount = float(budget.amount) if budget else 0.0
    spent_month = float(totals['month'])
    budget_left = budget_amount - spent_month
    budget_percent = round(min(spent_month / budget_amount * 100, 100), 1) if budget_amount else 0

    days_elapsed = today.day
    avg_per_day = round(spent_month / days_elapsed, 2) if days_elapsed else 0
    expense_count = qs.filter(date__gte=month_start).count()

    # ⭐ AI Insights
    insights = generate_insights(user, qs, today)

    # ⭐ Heatmap
    heatmap = build_heatmap(qs, today)

    # ⭐ Savings Goals
    goals = SavingsGoal.objects.filter(user=user)
    active_goals = [g for g in goals if not g.is_achieved][:3]

    # ⭐ Budget SMS alert check
    if budget and budget_amount > 0:
        try:
            check_budget_alerts(user, spent_month, budget_amount, CURRENCY_SYMBOLS.get('INR', '₹'))
        except Exception:
            pass

    # ⭐ Money Personality
    personality = analyze_personality(user, qs, today)

    context = {
        'totals': totals,
        'top_category': top_category,
        'avg_per_day': avg_per_day,
        'expense_count': expense_count,
        'recent_expenses': qs[:6],
        'category_labels': category_labels,
        'category_values': category_values,
        'category_colors': category_colors,
        'trend_labels': trend_labels,
        'trend_values': trend_values,
        'budget_amount': budget_amount,
        'budget_left': budget_left,
        'budget_percent': budget_percent,
        'has_budget': bool(budget),
        'today': today,
        'insights': insights,
        'heatmap': heatmap,
        'active_goals': active_goals,
        'personality': personality,
    }
    return render(request, 'expenses/dashboard.html', context)


# ============================================================
# EXPENSE CRUD
# ============================================================
@login_required
def expense_list(request):
    qs = Expense.objects.filter(user=request.user)
    qs, period = _apply_filters(request, qs)

    total = _total(qs)
    count = qs.count()

    paginator = Paginator(qs, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    params = request.GET.copy()
    params.pop('page', None)

    context = {
        'page_obj': page_obj,
        'total': total,
        'count': count,
        'period': period,
        'categories': CATEGORY_CHOICES,
        'querystring': params.urlencode(),
        'current_category': request.GET.get('category', ''),
        'current_q': request.GET.get('q', ''),
        'current_sort': request.GET.get('sort', '-date'),
        'start': request.GET.get('start', ''),
        'end': request.GET.get('end', ''),
    }
    return render(request, 'expenses/expense_list.html', context)


@login_required
def expense_create(request):
    if request.method == 'POST':
        form = ExpenseForm(request.POST, request.FILES)
        if form.is_valid():
            expense = form.save(commit=False)
            expense.user = request.user
            if not expense.category:
                expense.category, _conf = predict_category(expense.title)
            expense.save()
            messages.success(request, f'Expense added under "{expense.category}" ✅')
            return redirect('expense_list')
        messages.error(request, 'Please correct the errors below.')
    else:
        form = ExpenseForm(initial={'date': timezone.localdate()})

    return render(request, 'expenses/expense_form.html', {
        'form': form,
        'title': 'Add Expense',
        'subtitle': 'Log a new spending entry',
        'is_edit': False,
    })


@login_required
def expense_update(request, pk):
    expense = get_object_or_404(Expense, pk=pk, user=request.user)

    if request.method == 'POST':
        form = ExpenseForm(request.POST, request.FILES, instance=expense)
        if form.is_valid():
            obj = form.save(commit=False)
            if not obj.category:
                obj.category, _conf = predict_category(obj.title)
            obj.save()
            messages.success(request, 'Expense updated successfully ✏️')
            return redirect('expense_list')
        messages.error(request, 'Please correct the errors below.')
    else:
        form = ExpenseForm(instance=expense)

    return render(request, 'expenses/expense_form.html', {
        'form': form,
        'title': 'Edit Expense',
        'subtitle': f'Updating "{expense.title}"',
        'is_edit': True,
        'expense': expense,
    })


@login_required
def expense_delete(request, pk):
    expense = get_object_or_404(Expense, pk=pk, user=request.user)
    if request.method == 'POST':
        expense.delete()
        messages.success(request, 'Expense deleted 🗑️')
        return redirect('expense_list')
    return render(request, 'expenses/expense_confirm_delete.html', {'expense': expense})


# ============================================================
# API — Category Prediction (ML + Keyword)
# ============================================================
@login_required
def predict_category_api(request):
    text = request.GET.get('q', '')
    category, confidence, method = smart_predict(text)
    return JsonResponse({
        'category': category,
        'confidence': confidence,
        'method': method,
        'icon': CATEGORY_META.get(category, CATEGORY_META['Other'])['icon'],
        'color': CATEGORY_META.get(category, CATEGORY_META['Other'])['color'],
    })


# ============================================================
# BUDGET
# ============================================================
@login_required
def manage_budget(request):
    user = request.user
    today = timezone.localdate()
    month_start = today.replace(day=1)
    budget = Budget.objects.filter(user=user, month=month_start).first()

    if request.method == 'POST':
        form = BudgetForm(request.POST, instance=budget)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.user = user
            obj.month = month_start
            obj.save()
            messages.success(request, 'Monthly budget saved 🎯')
            return redirect('manage_budget')
    else:
        form = BudgetForm(instance=budget)

    spent = _total(Expense.objects.filter(user=user, date__gte=month_start, date__lte=today))
    budget_amount = budget.amount if budget else Decimal('0')
    remaining = budget_amount - spent
    percent = round(min(float(spent) / float(budget_amount) * 100, 100), 1) if budget_amount else 0

    history = []
    cursor = month_start
    for _ in range(6):
        b = Budget.objects.filter(user=user, month=cursor).first()
        s = _total(Expense.objects.filter(
            user=user,
            date__gte=cursor,
            date__lte=(cursor.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1),
        ))
        history.append({
            'month': cursor.strftime('%b %Y'),
            'budget': b.amount if b else None,
            'spent': s,
            'percent': round(min(float(s) / float(b.amount) * 100, 100), 1) if b and b.amount else 0,
        })
        cursor = (cursor - timedelta(days=1)).replace(day=1)

    return render(request, 'expenses/budget.html', {
        'form': form,
        'budget': budget,
        'spent': spent,
        'remaining': remaining,
        'percent': percent,
        'month_name': month_start.strftime('%B %Y'),
        'history': history,
    })


# ============================================================
# SAVINGS GOALS
# ============================================================
@login_required
def goals_list(request):
    user = request.user
    goals = SavingsGoal.objects.filter(user=user)

    total_target = sum(g.target_amount for g in goals) if goals else Decimal('0')
    total_saved = sum(g.saved_amount for g in goals) if goals else Decimal('0')
    overall_pct = round(float(total_saved / total_target * 100), 1) if total_target > 0 else 0

    return render(request, 'expenses/goals_list.html', {
        'goals': goals,
        'total_target': total_target,
        'total_saved': total_saved,
        'overall_pct': overall_pct,
    })


@login_required
def goal_create(request):
    if request.method == 'POST':
        form = SavingsGoalForm(request.POST)
        if form.is_valid():
            goal = form.save(commit=False)
            goal.user = request.user
            if not goal.icon:
                goal.icon = 'bi-trophy-fill'
            if not goal.color:
                goal.color = '#6366f1'
            goal.save()
            messages.success(request, f'Goal "{goal.name}" created! 🎯')
            return redirect('goals_list')
        messages.error(request, 'Please fix the errors.')
    else:
        form = SavingsGoalForm(initial={
            'saved_amount': 0,
            'icon': 'bi-trophy-fill',
            'color': '#6366f1',
        })

    return render(request, 'expenses/goal_form.html', {
        'form': form,
        'title': 'New Savings Goal',
        'is_edit': False,
    })


@login_required
def goal_update(request, pk):
    goal = get_object_or_404(SavingsGoal, pk=pk, user=request.user)

    if request.method == 'POST':
        form = SavingsGoalForm(request.POST, instance=goal)
        if form.is_valid():
            form.save()
            messages.success(request, 'Goal updated ✏️')
            return redirect('goals_list')
    else:
        form = SavingsGoalForm(instance=goal)

    return render(request, 'expenses/goal_form.html', {
        'form': form,
        'title': 'Edit Goal',
        'is_edit': True,
        'goal': goal,
    })


@login_required
def goal_delete(request, pk):
    goal = get_object_or_404(SavingsGoal, pk=pk, user=request.user)
    if request.method == 'POST':
        goal.delete()
        messages.success(request, 'Goal deleted 🗑️')
        return redirect('goals_list')
    return render(request, 'expenses/goal_confirm_delete.html', {'goal': goal})


@login_required
def goal_add_savings(request, pk):
    goal = get_object_or_404(SavingsGoal, pk=pk, user=request.user)
    if request.method == 'POST':
        try:
            amount = Decimal(request.POST.get('amount', '0'))
            if amount > 0:
                goal.saved_amount += amount
                goal.save()
                messages.success(request, f'₹{amount} added to "{goal.name}"! 💰')
            else:
                messages.error(request, 'Amount must be positive.')
        except Exception:
            messages.error(request, 'Invalid amount.')
    return redirect('goals_list')


# ============================================================
# RECURRING EXPENSES
# ============================================================
@login_required
def recurring_list(request):
    user = request.user
    items = RecurringExpense.objects.filter(user=user)
    due_count = items.filter(is_active=True, next_due__lte=timezone.localdate()).count()
    monthly_total = sum((r.monthly_equivalent for r in items if r.is_active), Decimal('0'))

    return render(request, 'expenses/recurring_list.html', {
        'items': items,
        'due_count': due_count,
        'monthly_total': monthly_total,
    })


@login_required
def recurring_create(request):
    if request.method == 'POST':
        form = RecurringExpenseForm(request.POST)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.user = request.user
            if not obj.icon:
                obj.icon = 'bi-arrow-repeat'
            if not obj.color:
                obj.color = '#6366f1'
            obj.save()
            messages.success(request, f'Recurring "{obj.title}" created! 🔁')
            return redirect('recurring_list')
    else:
        form = RecurringExpenseForm(initial={
            'start_date': timezone.localdate(),
            'next_due': timezone.localdate(),
            'frequency': 'monthly',
            'is_active': True,
            'icon': 'bi-arrow-repeat',
            'color': '#6366f1',
        })

    return render(request, 'expenses/recurring_form.html', {
        'form': form,
        'title': 'New Recurring Expense',
        'is_edit': False,
    })


@login_required
def recurring_update(request, pk):
    obj = get_object_or_404(RecurringExpense, pk=pk, user=request.user)

    if request.method == 'POST':
        form = RecurringExpenseForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            messages.success(request, 'Recurring updated ✏️')
            return redirect('recurring_list')
    else:
        form = RecurringExpenseForm(instance=obj)

    return render(request, 'expenses/recurring_form.html', {
        'form': form,
        'title': 'Edit Recurring',
        'is_edit': True,
        'item': obj,
    })


@login_required
def recurring_delete(request, pk):
    obj = get_object_or_404(RecurringExpense, pk=pk, user=request.user)
    if request.method == 'POST':
        obj.delete()
        messages.success(request, 'Recurring deleted 🗑️')
        return redirect('recurring_list')
    return render(request, 'expenses/recurring_confirm_delete.html', {'item': obj})


@login_required
def recurring_mark_paid(request, pk):
    obj = get_object_or_404(RecurringExpense, pk=pk, user=request.user)
    if request.method == 'POST':
        Expense.objects.create(
            user=request.user,
            title=obj.title,
            amount=obj.amount,
            currency='INR',
            category=obj.category,
            date=timezone.localdate(),
            notes=f'Auto-added from recurring ({obj.frequency})',
        )
        obj.last_paid = timezone.localdate()
        obj.advance_next_due()
        obj.save()
        messages.success(request, f'"{obj.title}" marked as paid ✅ Next: {obj.next_due}')
    return redirect('recurring_list')


@login_required
def recurring_toggle(request, pk):
    obj = get_object_or_404(RecurringExpense, pk=pk, user=request.user)
    if request.method == 'POST':
        obj.is_active = not obj.is_active
        obj.save()
        messages.success(request, f'"{obj.title}" {"activated" if obj.is_active else "paused"}!')
    return redirect('recurring_list')


# ============================================================
# SPLIT EXPENSES
# ============================================================
@login_required
def split_list(request):
    user = request.user
    splits = SplitExpense.objects.filter(user=user).prefetch_related('participants')

    total_receivable = Decimal('0')
    total_payable = Decimal('0')
    for s in splits:
        total_receivable += s.amount_to_receive
        total_payable += s.amount_to_pay

    return render(request, 'expenses/split_list.html', {
        'splits': splits,
        'total_receivable': total_receivable,
        'total_payable': total_payable,
    })


@login_required
def split_create(request):
    if request.method == 'POST':
        form = SplitExpenseForm(request.POST)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.user = request.user
            obj.save()
            formset = SplitParticipantFormSet(request.POST, instance=obj)
            if formset.is_valid():
                formset.save()
                messages.success(request, f'Split "{obj.title}" created! 👥')
                return redirect('split_detail', pk=obj.pk)
            else:
                obj.delete()
                messages.error(request, 'Please fix participant errors.')
    else:
        form = SplitExpenseForm(initial={'date': timezone.localdate(), 'paid_by_me': True})
        formset = SplitParticipantFormSet()

    return render(request, 'expenses/split_form.html', {
        'form': form,
        'formset': formset,
        'title': 'New Split Expense',
        'is_edit': False,
    })


@login_required
def split_update(request, pk):
    obj = get_object_or_404(SplitExpense, pk=pk, user=request.user)

    if request.method == 'POST':
        form = SplitExpenseForm(request.POST, instance=obj)
        formset = SplitParticipantFormSet(request.POST, instance=obj)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, 'Split updated ✏️')
            return redirect('split_detail', pk=obj.pk)
    else:
        form = SplitExpenseForm(instance=obj)
        formset = SplitParticipantFormSet(instance=obj)

    return render(request, 'expenses/split_form.html', {
        'form': form,
        'formset': formset,
        'title': 'Edit Split',
        'is_edit': True,
        'split': obj,
    })


@login_required
def split_detail(request, pk):
    obj = get_object_or_404(SplitExpense, pk=pk, user=request.user)
    return render(request, 'expenses/split_detail.html', {
        'split': obj,
        'participants': obj.participants.all(),
    })


@login_required
def split_delete(request, pk):
    obj = get_object_or_404(SplitExpense, pk=pk, user=request.user)
    if request.method == 'POST':
        obj.delete()
        messages.success(request, 'Split deleted 🗑️')
        return redirect('split_list')
    return render(request, 'expenses/split_confirm_delete.html', {'split': obj})


@login_required
def split_settle(request, pk, participant_pk):
    obj = get_object_or_404(SplitExpense, pk=pk, user=request.user)
    participant = get_object_or_404(SplitParticipant, pk=participant_pk, split=obj)
    if request.method == 'POST':
        participant.is_settled = True
        participant.settled_at = timezone.now()
        participant.save()
        messages.success(request, f'{participant.name} settled! ✅')
    return redirect('split_detail', pk=obj.pk)


# ============================================================
# GROUPS
# ============================================================
@login_required
def group_list(request):
    memberships = GroupMember.objects.filter(user=request.user).select_related('group')
    groups = [m.group for m in memberships]
    return render(request, 'expenses/group_list.html', {'groups': groups})


@login_required
def group_create(request):
    if request.method == 'POST':
        form = GroupForm(request.POST)
        if form.is_valid():
            group = form.save(commit=False)
            group.created_by = request.user
            if not group.icon:
                group.icon = 'bi-people-fill'
            if not group.color:
                group.color = '#6366f1'
            group.save()
            GroupMember.objects.create(group=group, user=request.user, role='owner')
            messages.success(request, f'Group "{group.name}" created! 🎉')
            return redirect('group_detail', pk=group.pk)
    else:
        form = GroupForm(initial={'icon': 'bi-people-fill', 'color': '#6366f1'})

    return render(request, 'expenses/group_form.html', {
        'form': form,
        'title': 'New Group',
        'is_edit': False,
    })


@login_required
def group_detail(request, pk):
    group = get_object_or_404(Group, pk=pk)
    membership = GroupMember.objects.filter(group=group, user=request.user).first()
    if not membership:
        messages.error(request, 'You are not a member of this group.')
        return redirect('group_list')

    members = GroupMember.objects.filter(group=group).select_related('user')
    expenses = group.expenses.select_related('paid_by')[:30]

    total = float(group.total_spent)
    member_totals = []
    for m in members:
        paid = float(group.expenses.filter(paid_by=m.user).aggregate(t=Sum('amount'))['t'] or 0)
        member_totals.append({
            'member': m,
            'paid': paid,
            'share': total / group.member_count if group.member_count else 0,
            'pct': (paid / total * 100) if total else 0,
        })

    return render(request, 'expenses/group_detail.html', {
        'group': group,
        'membership': membership,
        'members': members,
        'expenses': expenses,
        'member_totals': member_totals,
        'total': total,
        'invite_form': GroupInviteForm(),
    })


@login_required
def group_add_expense(request, pk):
    group = get_object_or_404(Group, pk=pk)
    if not GroupMember.objects.filter(group=group, user=request.user).exists():
        return redirect('group_list')

    if request.method == 'POST':
        form = GroupExpenseForm(request.POST)
        if form.is_valid():
            exp = form.save(commit=False)
            exp.group = group
            exp.paid_by = request.user
            exp.save()
            messages.success(request, f'₹{exp.amount} added to "{group.name}" ✅')
            return redirect('group_detail', pk=group.pk)
    else:
        form = GroupExpenseForm(initial={'date': timezone.localdate()})

    return render(request, 'expenses/group_expense_form.html', {
        'form': form,
        'group': group,
    })


@login_required
def group_invite(request, pk):
    group = get_object_or_404(Group, pk=pk)
    if not GroupMember.objects.filter(group=group, user=request.user).exists():
        return redirect('group_list')

    if request.method == 'POST':
        form = GroupInviteForm(request.POST)
        if form.is_valid():
            user = form.found_user
            if GroupMember.objects.filter(group=group, user=user).exists():
                messages.warning(request, f'{user.username} is already in this group.')
            else:
                GroupMember.objects.create(group=group, user=user, role='member')
                messages.success(request, f'{user.username} added! 👥')
            return redirect('group_detail', pk=group.pk)

    messages.error(request, 'Invalid invitation.')
    return redirect('group_detail', pk=group.pk)


@login_required
def group_leave(request, pk):
    group = get_object_or_404(Group, pk=pk)
    membership = GroupMember.objects.filter(group=group, user=request.user).first()
    if request.method == 'POST' and membership:
        if membership.role == 'owner' and group.member_count > 1:
            messages.error(request, 'Transfer ownership before leaving.')
        else:
            membership.delete()
            messages.success(request, f'You left "{group.name}".')
            return redirect('group_list')
    return redirect('group_detail', pk=group.pk)


@login_required
def group_delete(request, pk):
    group = get_object_or_404(Group, pk=pk)
    membership = GroupMember.objects.filter(group=group, user=request.user).first()
    if not membership or membership.role != 'owner':
        messages.error(request, 'Only the owner can delete this group.')
        return redirect('group_detail', pk=group.pk)

    if request.method == 'POST':
        name = group.name
        group.delete()
        messages.success(request, f'Group "{name}" deleted.')
        return redirect('group_list')

    return render(request, 'expenses/group_confirm_delete.html', {'group': group})


# ============================================================
# SETTINGS — Currency, SMS
# ============================================================
@login_required
def currency_settings(request):
    profile, _created = UserProfile.objects.get_or_create(user=request.user)

    if request.method == 'POST':
        form = UserProfileForm(request.POST, instance=profile)
        if form.is_valid():
            form.save()
            messages.success(request, 'Settings saved! 💱')
            return redirect('currency_settings')
    else:
        form = UserProfileForm(instance=profile)

    return render(request, 'expenses/currency_settings.html', {
        'form': form,
        'profile': profile,
        'rates': EXCHANGE_RATES_TO_INR,
        'symbols': CURRENCY_SYMBOLS,
    })


@login_required
def sms_settings(request):
    profile, _created = UserProfile.objects.get_or_create(user=request.user)

    if request.method == 'POST':
        profile.phone_number = request.POST.get('phone_number', '').strip()
        profile.sms_alerts_enabled = 'sms_alerts_enabled' in request.POST
        profile.save()
        messages.success(request, 'SMS settings saved! 📱')
        return redirect('sms_settings')

    return render(request, 'expenses/sms_settings.html', {'profile': profile})


@login_required
def sms_test(request):
    if request.method != 'POST':
        return redirect('sms_settings')

    profile = getattr(request.user, 'profile', None)
    if not profile or not profile.phone_number:
        messages.error(request, 'Please add a phone number first.')
        return redirect('sms_settings')

    ok, info = send_sms(
        profile.phone_number,
        f'✅ SpendWise test SMS — alerts working! '
        f'Sent at {timezone.now().strftime("%I:%M %p")}.'
    )
    if ok:
        messages.success(request, info)
    else:
        messages.error(request, info)
    return redirect('sms_settings')


# ============================================================
# PERSONALITY + ML
# ============================================================
@login_required
def personality_view(request):
    qs = Expense.objects.filter(user=request.user)
    personality = analyze_personality(request.user, qs)
    return render(request, 'expenses/personality.html', {'personality': personality})


@login_required
def ml_train_view(request):
    if not request.user.is_staff:
        messages.error(request, 'Only staff members can train the model.')
        return redirect('dashboard')

    info = model_info()

    if request.method == 'POST':
        result = train_model()
        if result['ok']:
            messages.success(request, result['message'])
        else:
            messages.warning(request, result['message'])
        return redirect('ml_train_view')

    return render(request, 'expenses/ml_train.html', {'info': info})


# ============================================================
# EXPORTS
# ============================================================
@login_required
def export_csv(request):
    qs = Expense.objects.filter(user=request.user)
    qs, period = _apply_filters(request, qs)

    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="expenses_{timezone.localdate()}.csv"'
    response.write('\ufeff')

    writer = csv.writer(response)
    writer.writerow(['Smart Expense Tracker — Expense Report'])
    writer.writerow([f'User: {request.user.username}', f'Generated: {timezone.localdate()}'])
    writer.writerow([])
    writer.writerow(['Date', 'Description', 'Category', 'Amount', 'Currency', 'Amount (INR)', 'Notes'])

    for e in qs:
        writer.writerow([
            e.date.strftime('%d-%m-%Y'), e.title, e.category,
            e.amount, e.currency, e.amount_in_inr, e.notes or ''
        ])

    writer.writerow([])
    writer.writerow(['', '', 'TOTAL', '', '', _total(qs), ''])
    return response


@login_required
def export_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    qs = Expense.objects.filter(user=request.user)
    qs, period = _apply_filters(request, qs)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="report_{timezone.localdate()}.pdf"'

    doc = SimpleDocTemplate(
        response, pagesize=A4,
        leftMargin=16 * mm, rightMargin=16 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('T', parent=styles['Title'], fontSize=22,
                                 textColor=colors.HexColor('#4f46e5'))
    sub_style = ParagraphStyle('S', parent=styles['Normal'], fontSize=10,
                               textColor=colors.HexColor('#64748b'))

    elements = [
        Paragraph('Smart Expense Tracker', title_style),
        Paragraph(f'Report &nbsp;|&nbsp; User: <b>{request.user.username}</b>', sub_style),
        Paragraph(f'{timezone.localdate().strftime("%d %b %Y")} &nbsp;|&nbsp; '
                  f'Records: {qs.count()}', sub_style),
        Spacer(1, 12),
    ]

    data = [['#', 'Date', 'Description', 'Category', 'Amount (INR)']]
    for idx, e in enumerate(qs, start=1):
        data.append([
            str(idx), e.date.strftime('%d-%m-%Y'),
            Paragraph(e.title[:55], styles['Normal']),
            e.category, f'{e.amount_in_inr:,.2f}',
        ])

    data.append(['', '', '', 'TOTAL', f'{_total(qs):,.2f}'])

    table = Table(data, colWidths=[10 * mm, 24 * mm, 72 * mm, 32 * mm, 32 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4f46e5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (4, 0), (4, -1), 'RIGHT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f1f5f9')]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e0e7ff')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#cbd5e1')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)
    doc.build(elements)
    return response


@login_required
def export_excel(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        messages.error(request, 'openpyxl not installed. Run: pip install openpyxl')
        return redirect('expense_list')

    qs = Expense.objects.filter(user=request.user)
    qs, period = _apply_filters(request, qs)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Expenses'

    title_font = Font(bold=True, size=18, color='FFFFFF')
    title_fill = PatternFill('solid', fgColor='4F46E5')
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='6366F1')
    total_fill = PatternFill('solid', fgColor='E0E7FF')
    total_font = Font(bold=True, size=12)
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    ws.merge_cells('A1:E1')
    ws['A1'] = 'Smart Expense Tracker — Report'
    ws['A1'].font = title_font
    ws['A1'].fill = title_fill
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    ws.merge_cells('A2:E2')
    ws['A2'] = f'User: {request.user.username}  |  Generated: {timezone.localdate()}'
    ws['A2'].alignment = Alignment(horizontal='center')
    ws['A2'].font = Font(italic=True, color='64748B')

    headers = ['#', 'Date', 'Description', 'Category', 'Amount (INR)']
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center')
        c.border = border

    for idx, e in enumerate(qs, start=1):
        row = idx + 4
        ws.cell(row=row, column=1, value=idx).border = border
        ws.cell(row=row, column=2, value=e.date.strftime('%d-%m-%Y')).border = border
        ws.cell(row=row, column=3, value=e.title).border = border
        ws.cell(row=row, column=4, value=e.category).border = border
        amt = ws.cell(row=row, column=5, value=float(e.amount_in_inr))
        amt.number_format = '₹#,##0.00'
        amt.border = border

    total_row = qs.count() + 5
    ws.cell(row=total_row, column=4, value='TOTAL').font = total_font
    ws.cell(row=total_row, column=4).fill = total_fill
    ws.cell(row=total_row, column=4).border = border
    tc = ws.cell(row=total_row, column=5, value=float(_total(qs)))
    tc.font = total_font
    tc.fill = total_fill
    tc.number_format = '₹#,##0.00'
    tc.border = border

    for col, width in zip('ABCDE', [6, 14, 42, 18, 16]):
        ws.column_dimensions[col].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="expenses_{timezone.localdate()}.xlsx"'
    return response


# ============================================================
# EMAIL
# ============================================================
@login_required
def send_summary_now(request):
    user = request.user
    if not user.email:
        messages.error(request, 'No email address on your account.')
        return redirect('dashboard')

    today = timezone.localdate()
    week_start = today - timedelta(days=7)

    qs = Expense.objects.filter(user=user, date__gte=week_start, date__lte=today)
    total = _total(qs)
    count = qs.count()

    cat_totals = qs.values('category').annotate(t=Sum('amount_in_inr')).order_by('-t')
    cat_lines = '\n'.join([f'  • {r["category"] or "Other"}: ₹{r["t"]:,.2f}' for r in cat_totals[:5]])

    subject = f'SpendWise Weekly Summary · {week_start} to {today}'
    body = f"""Hi {user.first_name or user.username},

Weekly spending summary:

📅 Period: {week_start} to {today}
💸 Total Spent: ₹{total:,.2f}
📊 Transactions: {count}

Top Categories:
{cat_lines or '  (no expenses)'}

Keep tracking!
— SpendWise
"""

    try:
        send_mail(subject, body, settings.DEFAULT_FROM_EMAIL, [user.email], fail_silently=False)
        messages.success(request, f'✅ Summary sent to {user.email}')
    except Exception as e:
        messages.error(request, f'Email failed: {e}')

    return redirect('dashboard')